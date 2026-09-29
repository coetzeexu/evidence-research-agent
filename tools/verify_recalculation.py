"""Recompute unchanged/modified Excel parameters in a separately supplied LibreOffice runtime."""

import argparse
import json
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from research_app.analytics import month_end_anchors, monthly_formula_backtest
from research_app.config import PROJECT_ROOT
from research_app.domain import ResearchBundle, now_iso

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--soffice", type=Path, required=True)
    parser.add_argument("--work-dir", type=Path, required=True)
    args = parser.parse_args()
    results = []
    for sample in ["nvda", "gold-bitcoin"]:
        root = PROJECT_ROOT / "samples" / sample
        bundle = ResearchBundle.model_validate_json((root / "bundle.json").read_text())
        for modified in [False, True]:
            spec = bundle.spec.model_copy(deep=True)
            edits = {}
            if modified:
                spec.weights = [1.0] if len(spec.symbols) == 1 else [0.7, 0.3]
                spec.cost_bps, spec.rebalance_months, spec.initial_capital = 25, 3, 50000
                edits = {
                    "B2": spec.initial_capital,
                    "B3": spec.cost_bps,
                    "B4": spec.rebalance_months,
                    **{f"B{8 + i}": weight for i, weight in enumerate(spec.weights)},
                }
            case = args.work_dir.resolve() / f"{sample}-{'modified' if modified else 'default'}"
            source = case / "input.xlsx"
            output = case / "recalculated"
            output.mkdir(parents=True, exist_ok=True)
            with ZipFile(root / "artifacts/report.xlsx") as original:
                names = load_workbook(root / "artifacts/report.xlsx", read_only=True).sheetnames
                parameter_path = f"xl/worksheets/sheet{names.index('Parameters') + 1}.xml"
                with ZipFile(source, "w", ZIP_DEFLATED) as changed:
                    for item in original.infolist():
                        content = original.read(item.filename)
                        if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                            xml = ET.fromstring(content)
                            # Direct OOXML parameter edits do not mark dependent cells dirty.
                            # Discard delivery caches so this verifies fresh spreadsheet calculation.
                            for cell in xml.findall(".//s:c", NS):
                                if cell.find("s:f", NS) is not None:
                                    cached = cell.find("s:v", NS)
                                    if cached is not None:
                                        cell.remove(cached)
                            if modified and item.filename == parameter_path:
                                for address, value in edits.items():
                                    cell = xml.find(f'.//s:c[@r="{address}"]', NS)
                                    assert cell is not None, address
                                    cell.attrib.pop("t", None)
                                    cell.find("s:v", NS).text = str(value)
                            content = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
                        changed.writestr(item, content)
            process = subprocess.run(
                [
                    str(args.soffice.resolve()),
                    f"-env:UserInstallation={(case / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    str(output),
                    str(source),
                ],
                capture_output=True,
                text=True,
                timeout=90,
            )
            assert process.returncode == 0 and (output / source.name).exists(), (
                "LibreOffice conversion failed"
            )
            book = load_workbook(output / source.name, data_only=True)
            errors = [(s.title, c.coordinate) for s in book for row in s for c in row if c.data_type == "e"]
            assert not errors, errors
            anchors = (
                bundle.comparison.get("anchors", [])
                if len(spec.symbols) > 1
                else month_end_anchors(bundle.datasets, spec.symbols, spec.start, spec.end)
            )
            expected = monthly_formula_backtest(anchors, spec)
            nav_col = len(spec.symbols) + 6
            diffs = [
                abs(book["MonthlyCalc"].cell(i + 2, nav_col).value - row["nav"])
                for i, row in enumerate(expected["rows"])
            ]
            assert max(diffs, default=0) < 1e-6, max(diffs)
            results.append(
                {
                    "sample": sample,
                    "modified": modified,
                    "passed": True,
                    "monthly_rows": len(diffs),
                    "max_nav_difference_usd": max(diffs, default=0),
                    "formula_errors": errors,
                }
            )
    report = {"checked_at": now_iso(), "checks": results}
    (PROJECT_ROOT / "evals/excel-recalculation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    )
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
