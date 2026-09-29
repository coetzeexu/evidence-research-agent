from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .domain import ResearchBundle
from .report_content import decisions, key_points, limitations, percent
from .security import allowed_link


def excerpt(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    end = text.rfind("。", 0, limit)
    if end < 45:
        end = text.rfind("；", 0, limit)
    return (text[: end + 1] if end >= 45 else text[:limit] + "…") + "完整内容见 HTML。"


def hyperlink(paragraph, label: str, url: str):
    if not allowed_link(url):
        paragraph.add_run(label)
        return
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "47765C")
    props.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    props.append(size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def export_document(bundle: ResearchBundle, target: Path, charts: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.7)
    sec.left_margin, sec.right_margin = Inches(0.8), Inches(0.8)
    styles = doc.styles
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        style = styles[name]
        style.font.name = "Heiti SC"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if "Theme" in attr:
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), "Heiti SC")
        borders = style.element.find("w:pPr/w:pBdr", namespaces=style.element.nsmap)
        if borders is not None:
            borders.getparent().remove(borders)
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Title"].font.size = Pt(27)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.author = "Evidence"
    doc.core_properties.title = bundle.spec.title

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        if widths:
            for c, w in zip(t.columns, widths, strict=True):
                c.width = Inches(w)
        for i, value in enumerate(headers):
            t.rows[0].cells[i].text = str(value)
        for values in rows:
            for cell, value in zip(t.add_row().cells, values, strict=True):
                cell.text = str(value)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for r, row in enumerate(t.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ["top", "left", "bottom", "right"]:
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "4")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                tcPr.append(borders)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "3C5147" if r == 0 else "F2F5F1" if r % 2 == 0 else "FFFFFF")
                tcPr.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ["top", "bottom", "left", "right"]:
                    e = OxmlElement(f"w:{side}")
                    e.set(qn("w:w"), "95")
                    e.set(qn("w:type"), "dxa")
                    margins.append(e)
                tcPr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.2
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        if r == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    doc.add_paragraph("EVIDENCE / INVESTMENT RESEARCH", "Subtitle")
    doc.add_paragraph("策略研究报告", "Title")
    doc.add_paragraph(bundle.spec.title, "Heading 2")
    doc.add_paragraph(
        f"研究区间 {bundle.spec.start} 至 {bundle.spec.end}    生成 {bundle.created_at[:10]}\n研究版本 {bundle.id}    美元计价",
        "Subtitle",
    )
    doc.add_heading("研究结果", 1)
    doc.add_paragraph(
        "需求覆盖检查" + ("通过。" if bundle.quality.get("passed") else "尚有缺口，当前为部分交付。")
    )
    for requirement in bundle.quality.get("requirements", []):
        doc.add_paragraph(
            requirement["requirement"]
            + "："
            + (
                "已对应事件 " + ", ".join(requirement["event_ids"])
                if requirement["event_ids"]
                else "尚未核实"
            )
        )
    for point in key_points(bundle)[:2]:
        doc.add_paragraph(point)
    image = charts / ("comparison.png" if bundle.comparison.get("series") else "market.png")
    if image.exists():
        doc.add_picture(str(image), width=Inches(6.6))
    doc.add_paragraph(
        "图 1 共同月度净值（初始为 100）"
        if bundle.comparison.get("series")
        else "图 1 日线与已核验事件，交互细节见 HTML",
        "Caption",
    )
    metrics = bundle.comparison.get("asset_metrics") or [
        m for m in bundle.metrics if m["id"].startswith("performance-")
    ]
    table(
        ["资产", "累计收益", "年化收益", "年化波动", "最大回撤"],
        [
            [
                m["symbol"],
                percent(m["total_return"]),
                percent(m["cagr"]),
                percent(m["volatility"]),
                percent(m["max_drawdown"]),
            ]
            for m in metrics
        ],
    )
    doc.add_paragraph(
        (
            "比较表采用共同月度观测。"
            if bundle.comparison.get("asset_metrics")
            else "表中为各资产自身日度回顾，实际区间可能不同，不作共同月度比较。"
        )
        + "数据源与完整日线保存在 Excel 底稿，图表与数字来自同一研究快照。"
    )

    doc.add_heading("决策框架", 1).paragraph_format.page_break_before = True
    table(["目标", "待回答问题", "检验指标", "适用条件"], decisions(bundle), [0.65, 1.35, 1.75, 2.85])
    if bundle.comparison:
        doc.add_heading("压力期与通胀", 2)
        for point in key_points(bundle)[len(bundle.claims) :]:
            doc.add_paragraph(point)
        rows = bundle.comparison.get("inflation_summary", [])
        table(
            ["资产", "CPI 同比 ≥3% 月数", "该组月均收益", "CPI 同比 <3% 月数", "该组月均收益"],
            [
                [
                    r["symbol"],
                    r["high_inflation_n"],
                    percent(r["high_inflation_mean_return"]),
                    r["low_inflation_n"],
                    percent(r["low_inflation_mean_return"]),
                ]
                for r in rows
            ],
        )
        for disclosure in bundle.disclosures:
            doc.add_heading(disclosure["title"], 2)
            doc.add_paragraph(disclosure["description"] + disclosure["reason"])
            doc.add_paragraph(disclosure["limitations"])
            hyperlink(doc.add_paragraph(), "产品定义、费用与基准：GLD 官方资料", disclosure["url"])
    else:
        doc.add_heading("等级说明", 2)
        doc.add_paragraph(
            "市场反应强度为窗口相对基准收益的绝对值，除以事前日收益波动率乘以窗口天数平方根。达到 2 为高，达到 1 为中，其余为低。该指标是描述性标准化反应，不是显著性检验或因果概率。"
        )
        doc.add_paragraph(
            "关联可信度综合已读来源、事件日期精度和研究核验。没有精确发布时间的事件最高显示中等关联可信度，避免将日线噪声当作精确事件效应。"
        )

    doc.add_heading("方法与配置检验", 1).paragraph_format.page_break_before = True
    if bundle.comparison.get("backtest", {}).get("rows"):
        p = bundle.comparison["backtest"]["metrics"]
        a = bundle.comparison["formula_backtest"]["metrics"]
        doc.add_paragraph(
            f"初始资金 ${bundle.spec.initial_capital:,.0f}，权重 "
            + " / ".join(
                f"{s} {w:.0%}" for s, w in zip(bundle.spec.symbols, bundle.spec.weights, strict=True)
            )
            + f"，交易成本 {bundle.spec.cost_bps:g} bps，再平衡间隔 {bundle.spec.rebalance_months} 个月（0 为持有）。"
        )
        table(
            ["口径", "累计收益", "最大回撤"],
            [
                ["下一可交易开盘执行", percent(p["total_return"]), percent(p["max_drawdown"])],
                ["Excel 月度配置近似", percent(a["total_return"]), percent(a["max_drawdown"])],
            ],
        )
        doc.add_picture(str(charts / "portfolio.png"), width=Inches(6.6))
        doc.add_paragraph(
            "图 2 实际执行模拟与月度公式近似分别计算。调仓份额使用信号时点已知价格，执行价格取下一完整日线开盘；现金不足时限制买入。",
            "Caption",
        )
    sensitivity = bundle.comparison.get("sensitivity", {})
    if sensitivity:
        doc.add_heading("配置敏感性分析", 1).paragraph_format.page_break_before = True
        doc.add_paragraph(sensitivity["method"])
        table(
            ["情景", "年化收益", "最大回撤", "年化波动"],
            [
                [r["label"], percent(r["cagr"]), percent(r["max_drawdown"]), percent(r["volatility"])]
                for r in sensitivity["rows"]
            ],
            [2.3, 1.4, 1.4, 1.4],
        )
        doc.add_paragraph(sensitivity["caveat"])
        for r in sensitivity["rows"]:
            if r["dimension"] == "period":
                doc.add_paragraph(
                    f"{r['label']}：{r['start']} 至 {r['end']}，{r['observations']} 个月末锚点。"
                )
    for line in limitations(bundle)[:5]:
        doc.add_paragraph(line)
    for gap in bundle.quality.get("repair_actions", []):
        doc.add_paragraph(f"未解决要求：{gap['query']}；{gap['reason']}。")
    summary = bundle.review.get("summary", "数值与引用完整性检查已执行。")
    doc.add_paragraph(
        f"独立核验：{summary[:220]}" + ("。完整记录见 HTML 方法页。" if len(summary) > 220 else "")
    )

    sources = {s.id: s for s in bundle.sources}
    for start in range(0, len(bundle.events), 3):
        doc.add_heading(
            "事件与市场反应" + (f"（续 {start // 3}）" if start else ""), 1
        ).paragraph_format.page_break_before = True
        for event in bundle.events[start : start + 3]:
            doc.add_heading(f"{event.date}  {event.title}", 2)
            doc.add_paragraph(excerpt(event.summary, 180))
            annotations = [a for a in bundle.annotations if a["event_id"] == event.id]
            for annotation in annotations:
                windows = "；".join(
                    f"{w['days']} 日相对收益 {percent(w.get('relative_return'))}"
                    for w in annotation["windows"]
                )
                doc.add_paragraph(
                    f"{annotation['symbol']} 对齐交易日 {annotation['date']}，反应等级 {annotation['rating']}。{windows}。"
                )
            if event.uncertainty:
                p = doc.add_paragraph("时间与证据：" + excerpt(event.uncertainty, 130))
                for run in p.runs:
                    run.font.size = Pt(8)
            p = doc.add_paragraph()
            for i in event.source_ids:
                if i in sources:
                    hyperlink(p, f"[{i}] {sources[i].publisher}  ", sources[i].url)

    doc.add_heading("来源索引与复核", 1).paragraph_format.page_break_before = True
    doc.add_paragraph(
        "点击下列名称进入原始来源。数据快照 ID 与哈希可在 Excel Sources、HTML 数据页及 manifest.json 中交叉核对。报告仅保留短引用片段，完整内容请查阅原文。"
    )
    for source in bundle.sources:
        p = doc.add_paragraph()
        hyperlink(p, f"[{source.id}] {source.title}", source.url)
        r = p.add_run(f"\n采集 {source.retrieved_at[:19]} UTC；状态 {source.status}")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("879387")
    for symbol, dataset in bundle.datasets.items():
        p = doc.add_paragraph()
        hyperlink(p, f"[{dataset.id}] {symbol} 日线数据", dataset.source_url)
    doc.save(target)
