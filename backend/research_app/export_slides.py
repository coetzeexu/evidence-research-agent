from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from .analytics import month_end_anchors
from .domain import ResearchBundle
from .report_content import decisions, key_points, percent
from .security import allowed_link


def export_slides(bundle: ResearchBundle, target: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.author = "Evidence"
    prs.core_properties.title = bundle.spec.title
    colors = ["527B64", "C09A5B", "728CA7", "B58086", "8F81A9"]
    source_notes = "\n".join(f"{s.id}: {s.url}" for s in bundle.sources)
    source_notes += "\n" + "\n".join(f"{d.id}: {d.source_url}" for d in bundle.datasets.values())

    def text(slide, value, x, y, w, h, size=18, color="334A3D", bold=False, url=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        for i, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(9)
            run = p.add_run()
            run.text = line
            run.font.name = "Heiti SC"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            if url and allowed_link(url):
                run.hyperlink.address = url
        return box

    def slide(title, subtitle=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("FAFBF7")
        text(s, title, 0.7, 0.52, 11.9, 0.6, 28, bold=True)
        if subtitle:
            text(s, subtitle, 0.72, 1.19, 11.8, 0.42, 12, "788876")
        text(s, f"Evidence  /  {bundle.id}", 0.7, 7.06, 10, 0.22, 9, "B0B9A4")
        text(s, f"{len(prs.slides):02}", 12.05, 7.04, 0.5, 0.25, 10, "9BA88C")
        s.notes_slide.notes_text_frame.text = (
            f"研究版本 {bundle.id}\n{bundle.review.get('summary', '')}\n来源：\n{source_notes}"
        )
        return s

    def table(s, headers, rows, x=0.72, y=1.95, w=11.9, h=3.5, widths=None, size=15):
        t = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
        if widths:
            for c, width in zip(t.columns, widths, strict=True):
                c.width = Inches(width)
        for ri, row in enumerate([headers, *rows]):
            for ci, value in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = str(value)
                cell.margin_left = Inches(0.14)
                cell.margin_right = Inches(0.14)
                cell.margin_top = Inches(0.12)
                cell.margin_bottom = Inches(0.1)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(
                    "415E4E" if ri == 0 else "EEF2E8" if ri % 2 else "FFFFFF"
                )
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name = "Heiti SC"
                        r.font.size = Pt(size)
                        r.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "607853")
        return t

    def chart(s, dates, series, x=0.7, y=1.9, w=11.9, h=4.5, kind=XL_CHART_TYPE.LINE):
        data = CategoryChartData()
        data.categories = [d[:7] for d in dates]
        for name, values in series:
            data.add_series(name, values)
        c = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        c.has_legend = True
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
        c.legend.font.size = Pt(11)
        c.chart_style = 10
        for i, line in enumerate(c.series):
            line.format.line.color.rgb = RGBColor.from_string(colors[i % len(colors)])
            line.format.line.width = Pt(2.2)
        c.category_axis.tick_labels.font.size = Pt(9)
        c.value_axis.tick_labels.font.size = Pt(10)
        return c

    s = slide(
        "研究判断与证据", f"{' / '.join(bundle.spec.symbols)}    {bundle.spec.start} 至 {bundle.spec.end}"
    )
    text(s, bundle.spec.title, 0.75, 1.95, 11.7, 1.25, 29, "48603D", True)
    for i, point in enumerate(key_points(bundle)[:2]):
        text(s, point, 0.8, 3.7 + i * 0.9, 11.6, 0.8, 20, "617B65")
    text(
        s,
        "点击来源回链复核，原始数据和公式保留在随附 Excel，交互图表保留在 HTML。",
        0.8,
        6.15,
        11.6,
        0.5,
        14,
        "617B65",
    )

    s = slide("价格表现与共同估值", "原生可编辑图表。初始净值 100，所有资产采用同一估值锚点。")
    if bundle.comparison.get("series"):
        series = bundle.comparison["series"]
        chart(s, series[0]["dates"], [(r["symbol"], r["normalized"]) for r in series])
    elif bundle.comparison.get("available") is not False:
        anchors = month_end_anchors(bundle.datasets, bundle.spec.symbols, bundle.spec.start, bundle.spec.end)
        if anchors:
            chart(
                s,
                [r["date"] for r in anchors],
                [
                    (symbol, [r["prices"][symbol] / anchors[0]["prices"][symbol] * 100 for r in anchors])
                    for symbol in bundle.spec.symbols
                ],
            )
    else:
        text(s, "共同完整月份不足或存在缺失，暂不展示月度比较。", 0.8, 2.5, 11.5, 1.2, 22)

    metrics = bundle.comparison.get("asset_metrics") or [
        m for m in bundle.metrics if m["id"].startswith("performance-")
    ]
    s = slide(
        "收益与风险",
        "共同月度观测。"
        if bundle.comparison.get("asset_metrics")
        else "各资产自身日度回顾，实际区间可能不同，不作共同月度比较。",
    )
    table(
        s,
        ["资产", "累计收益", "年化收益", "年化波动", "最大回撤"],
        [
            [
                m["symbol"],
                percent(m["total_return"]),
                percent(m["cagr"]),
                percent(m["volatility"]),
                percent(m["max_drawdown"]),
            ]
            for m in metrics
        ],
        h=min(3.1, 0.65 * (len(metrics) + 1)),
        size=18,
    )
    row = bundle.comparison.get("backtest", {}).get("metrics", {})
    if row:
        text(s, "固定权重配置，下一开盘执行", 0.8, 4.15, 11.5, 0.5, 20, "5C784D", True)
        text(
            s,
            f"累计收益 {percent(row.get('total_return'))}      最大回撤 {percent(row.get('max_drawdown'))}      成本 {bundle.spec.cost_bps:g} bps",
            0.8,
            4.85,
            11.5,
            0.6,
            20,
            "617B65",
        )
        text(
            s,
            "Excel 月度近似可修改参数。下一开盘执行保留独立成交台账，两个结果不能混用。",
            0.8,
            5.9,
            11.5,
            0.65,
            16,
            "617B65",
        )
    else:
        text(
            s,
            "累计收益描述这段历史。判断事件影响需要同时检查基准收益、发布时点与同期其他事件。",
            0.8,
            4.4,
            11.4,
            1.1,
            22,
            "617B65",
        )

    q = bundle.quality
    s = slide("研究要求覆盖", "覆盖检查与事实核验分开；候选关联不构成因果证明。")
    text(s, "覆盖检查通过" if q.get("passed") else "执行已结束，研究仍有缺口", 0.8, 1.9, 11.5, 0.6, 24)
    items = [
        r["requirement"] + ("：已对应来源事件" if r["event_ids"] else "：尚未核实")
        for r in q.get("requirements", [])
    ]
    items += [a["query"] + "：" + a["reason"] for a in q.get("repair_actions", [])]
    if not items:
        items = [
            f"{len(bundle.events)} 个来源事件；{q.get('associated_changes', 0)} 个异动存在时间候选关联。"
        ]
    for i, item in enumerate(items[:6]):
        text(s, item, 0.8, 2.8 + i * 0.55, 11.5, 0.5, 15)
    if len(items) > 6:
        text(s, "其余缺口见 HTML 覆盖清单和 Excel Coverage。", 0.8, 6.25, 11.5, 0.4, 13)
    sensitivity = bundle.comparison.get("sensitivity", {})
    for start in range(0, len(sensitivity.get("rows", [])), 6):
        s = slide("配置敏感性分析", "固定网格逐项改变一个参数，采用下一开盘执行；不选择全样本最优配置。")
        table(
            s,
            ["情景", "区间", "年化收益", "最大回撤"],
            [
                [r["label"], f"{r['start']} / {r['end']}", percent(r["cagr"]), percent(r["max_drawdown"])]
                for r in sensitivity["rows"][start : start + 6]
            ],
            h=3.8,
            widths=[2.7, 4.2, 2.5, 2.5],
            size=14,
        )
        text(s, "分段是历史稳健性检查，不是样本外检验。Excel 敏感性表为快照。", 0.8, 6.15, 11.5, 0.5, 15)
    s = slide("决策框架", "研究目标决定检验方法，历史结果需要放回适用情景。")
    table(
        s,
        ["目标", "待回答问题", "检验指标", "适用条件"],
        decisions(bundle),
        h=4.35,
        widths=[0.9, 2.4, 3.5, 5.1],
        size=14,
    )

    if bundle.comparison:
        s = slide(
            "压力情景与实际购买力", f"全样本事后选取 {bundle.spec.benchmark} 最差月份，CPI 为当前历史版本。"
        )
        stress = bundle.comparison.get("stress_periods", [])[:5]
        table(
            s,
            ["月份", bundle.spec.benchmark, *bundle.spec.symbols],
            [
                [
                    r["date"][:7],
                    percent(r["benchmark_return"]),
                    *[percent(r["returns"][sym]) for sym in bundle.spec.symbols],
                ]
                for r in stress
            ],
            h=3.2,
            size=17,
        )
        text(
            s,
            "避险表现依赖压力类型。通胀检验应同时观察 CPI 调整净值、高通胀时期和持有期限。",
            0.8,
            5.65,
            11.6,
            0.8,
            19,
            "617B65",
        )
    for start in range(0, min(len(bundle.events), 12), 4):
        s = slide(
            "关键事件与证据" + (f" {start // 4 + 1}" if start else ""),
            "日期为来源支持的事件日；日线反应窗口保留时间精度限制。",
        )
        for i, event in enumerate(bundle.events[start : start + 4]):
            y = 1.94 + i * 1.15
            text(s, str(event.date), 0.78, y, 1.5, 0.4, 15, "617B65")
            title = event.title if len(event.title) < 51 else event.title[:48] + "…"
            source = next((r for r in bundle.sources if r.id in event.source_ids), None)
            text(s, title, 2.35, y, 9.6, 0.63, 19, "567148", True, source.url if source else None)
            reactions = []
            for a in (r for r in bundle.annotations if r["event_id"] == event.id):
                w = next((r for r in a["windows"] if r["days"] == 5), None)
                reactions.append(
                    f"{a['symbol']} {percent(w.get('relative_return') if w else None)}（{a['rating']}）"
                )
            text(
                s,
                f"5 日相对收益 {' · '.join(reactions)}    来源 {source.publisher if source else '详见索引'}",
                2.35,
                y + 0.64,
                9.6,
                0.38,
                12,
                "617B65",
            )
    if bundle.disclosures:
        s = slide("黄金的可投资代理：GLD", "产品定义、采用原因与回测口径")
        d = bundle.disclosures[0]
        for i, (heading, value) in enumerate(
            [("产品定义", d["description"]), ("采用原因", d["reason"]), ("适用边界", d["limitations"])]
        ):
            text(s, heading, 0.8, 2 + i * 1.35, 1.5, 0.5, 17, "759654", True)
            text(s, value, 2.55, 2 + i * 1.35, 9.7, 1.1, 18, "617B65")
        text(s, "SPDR Gold Shares 官方产品资料", 0.8, 6.25, 11.6, 0.45, 14, "6F9950", url=d["url"])

    s = slide(
        "方法与复核边界",
        "核验"
        + ("通过" if bundle.review.get("passed") else "未通过")
        + "；完整问题与来源保留在 Word、HTML 和演讲者备注。",
    )
    rows = [
        ("行情与事件", "保存完整日线、公司行动、发布日与数据快照哈希。日期级资料无法分离盘中反应。"),
        ("关联评级", "相对基准反应按事前波动标准化。高反应与高证据可信度是不同维度，均不证明因果。"),
        ("回测边界", "参数事先固定。宏观数据采用当前历史版本，只用于解释。跨市场日线时间和费用分别保留。"),
        ("可复核交付", "Excel 保留公式与台账，Word 保留方法与完整来源，HTML 支持缩放、事件筛选与来源跳转。"),
    ]
    for i, (heading, value) in enumerate(rows):
        text(s, heading, 0.8, 1.95 + i * 1.13, 2, 0.45, 18, "6E8F50", True)
        text(s, value, 3.0, 1.95 + i * 1.13, 9.5, 0.92, 16, "617B65")
    for start in range(0, len(bundle.sources), 10):
        s = slide("原始来源索引", "点击名称可进入原文。所有页面的演讲者备注也保留来源 URL。")
        for i, source in enumerate(bundle.sources[start : start + 10]):
            title = source.title if len(source.title) <= 87 else source.title[:86] + "…"
            label = f"{start + i + 1:02}  {title}"
            text(s, label, 0.8, 1.88 + i * 0.46, 11.7, 0.42, 13, "617B65", url=source.url)
    prs.save(target)
