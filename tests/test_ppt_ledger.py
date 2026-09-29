import hashlib
import json
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from research_app.export_slides import PPT_SKILL, _wrap, export_slides


def _text(prs):
    result = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                result.append(shape.text)
            if shape.has_table:
                result.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(result)


def test_ppt_uses_original_method_and_native_evidence(bundle, tmp_path):
    target = tmp_path / "report.pptx"
    export_slides(bundle, target)
    prs = Presentation(target)
    text = _text(prs)
    assert any("33%" in line for line in _wrap("资产涨幅约为 33%", 1.2, 18))
    assert "operating-review" in prs.core_properties.subject
    assert "ppt169_apple_fy2025_review" in prs.core_properties.keywords
    assert "同口径资产差异" in text and "原因尚未确立" in text
    assert text.index("收益与风险") < text.index("同口径资产差异") < text.index("原因尚未确立")
    assert "研究要求覆盖" not in text
    assert not any(t in text for t in ["随附 Excel", "HTML", "Word", "研究产物"])
    assert all(str(s.background.fill.fore_color.rgb) == "FFFFFF" for s in prs.slides)
    tables = [shape.table for s in prs.slides for shape in s.shapes if shape.has_table]
    assert tables
    numeric_cell = tables[0].cell(1, 1)
    assert numeric_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert numeric_cell.text_frame.paragraphs[0].runs[0].font.name == "Consolas"
    with ZipFile(target) as z:
        assert any(n.startswith("ppt/charts/chart") for n in z.namelist())
        assert any(n.startswith("ppt/embeddings/") for n in z.namelist())
        links = "\n".join(z.read(n).decode() for n in z.namelist() if n.endswith(".rels"))
        assert "https://example.com/source" in links
        assert "https://example.com/GLD" in links


def test_ppt_paginates_all_sensitivity_rows_and_preserves_sources(bundle, tmp_path):
    rows = bundle.comparison["sensitivity"]["rows"]
    assert len(rows) > 6
    target = tmp_path / "full.pptx"
    export_slides(bundle, target)
    prs = Presentation(target)
    text = _text(prs)
    assert all(r["label"] in text for r in rows)
    assert sum("配置敏感性" in s.shapes[0].text for s in prs.slides if s.shapes[0].has_text_frame) > 1
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
    assert bundle.datasets["GLD"].content_hash in notes
    assert "comparison.sensitivity.rows" in notes
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
            if shape.has_table:
                assert shape.top + shape.height < 6.78 * 914400


def test_ppt_untrusted_source_scheme_is_not_clickable(bundle, tmp_path):
    bundle.sources[0].url = "javascript:alert(1)"
    target = tmp_path / "safe.pptx"
    export_slides(bundle, target)
    with ZipFile(target) as z:
        relations = "".join(z.read(n).decode() for n in z.namelist() if n.endswith(".rels"))
        assert "javascript:" not in relations


def test_ppt_original_template_files_match_pinned_sources():
    manifest = json.loads((PPT_SKILL / "source-manifest.json").read_text())
    for entry in manifest["sources"]:
        path = PPT_SKILL / entry["local_path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == entry["sha256"]
    method = (PPT_SKILL / "references/operating-review.md").read_text()
    assert "style_id: operating-review" in method
    assert "Native Editability" in method and "Claim Discipline" in method


def test_ppt_requires_local_method_template(bundle, tmp_path, monkeypatch):
    from research_app import export_slides as module

    monkeypatch.setattr(module, "PPT_SKILL", tmp_path / "missing")
    with pytest.raises(FileNotFoundError):
        module.export_slides(bundle, tmp_path / "report.pptx")
