import base64
import hashlib
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from research_app.config import Settings
from research_app.export_document import export_document
from research_app.export_slides import export_slides
from research_app.export_workbook import export_workbook
from research_app.storage import Store


def test_workbook_real_formulas_cached_parity_and_source_links(bundle, tmp_path):
    target = tmp_path / "report.xlsx"
    export_workbook(bundle, target)
    formula = load_workbook(target, data_only=False)
    values = load_workbook(target, data_only=True)
    assert formula["MonthlyCalc"]["H2"].data_type == "f"
    assert values["Overview"]["B17"].value == bundle.comparison["formula_backtest"]["metrics"]["total_return"]
    assert formula["Sources"]["C2"].hyperlink.target == "https://example.com/source"
    for sheet in values:
        for row in sheet:
            assert not any(cell.data_type == "e" for cell in row)


def test_docx_and_native_editable_pptx(bundle, tmp_path):
    # No chart raster necessary for this structural Word check; visual QA uses generated samples.
    bundle.comparison = {}
    export_document(bundle, tmp_path / "report.docx", tmp_path)
    export_slides(bundle, tmp_path / "report.pptx")
    with ZipFile(tmp_path / "report.docx") as z:
        assert b"https://example.com/source" in z.read("word/_rels/document.xml.rels")
    with ZipFile(tmp_path / "report.pptx") as z:
        assert any(n.startswith("ppt/charts/chart") and n.endswith(".xml") for n in z.namelist())
        assert any(n.startswith("ppt/embeddings/") for n in z.namelist())


def test_store_recovery_and_refresh_flag(tmp_path):
    store = Store(tmp_path)
    rid = store.create("research", refresh=True)
    store.update(rid, status="running")
    store.recover()
    assert store.get(rid)["status"] == "queued"
    assert store.get(rid)["refresh_requested"] == 1
    store.emit(rid, "step", "collect")
    event = store.events(rid)[0]
    assert store.events(rid, event["seq"]) == []


def test_api_origin_paths_and_missing_configuration(monkeypatch, tmp_path, bundle):
    from research_app import api

    local = Store(tmp_path)
    monkeypatch.setattr(api, "store", local)
    monkeypatch.setattr(api, "config", Settings(tmp_path, "", "", "https://example.com"))
    client = TestClient(api.app)
    assert client.get("/api/health").json()["model_configured"] is False
    assert client.post("/api/runs", json={"prompt": "test"}).status_code == 409
    assert (
        client.post(
            "/api/runs", json={"prompt": "test"}, headers={"Origin": "https://evil.example"}
        ).status_code
        == 403
    )
    rid = local.create("test", run_id=bundle.id)
    local.save_json(local.run_dir(rid) / "bundle.json", bundle.model_dump(mode="json"))
    local.update(rid, bundle_path=str(local.run_dir(rid) / "bundle.json"), status="complete")
    assert client.get(f"/api/runs/{rid}/bundle").json()["id"] == bundle.id
    assert "bundle_path" not in client.get(f"/api/runs/{rid}").json()
    assert client.get(f"/api/runs/{rid}/artifacts/.env").status_code == 404
    assert "private-test-token" not in client.get("/api/health").text


def test_offline_html_embeds_data_safely_and_pins_script_hash(bundle, tmp_path, monkeypatch):
    from research_app import exporters

    assets = tmp_path / "dist" / "report"
    assets.mkdir(parents=True)
    script = "document.getElementById('root').textContent='ready'"
    (assets / "report.js").write_text(script)
    (assets / "report.css").write_text("body{color:#263d2e}")
    (tmp_path / "THIRD_PARTY_NOTICES.txt").write_text("MIT License · Copyright (c) 2026 Shane Levine")
    monkeypatch.setattr(exporters, "PROJECT_ROOT", tmp_path)
    bundle.events[0].summary = '</script><script src="https://evil.example/payload"></script>'
    target = tmp_path / "report.html"
    exporters.export_html(bundle, target)
    html = target.read_text()
    signature = base64.b64encode(hashlib.sha256(script.encode()).digest()).decode()
    assert f"script-src 'sha256-{signature}'" in html
    assert "connect-src 'none'" in html
    assert "Copyright (c) 2026 Shane Levine" in html
    assert '<script src="https://evil.example/payload">' not in html
    assert "\\u003c/script>" in html


def test_sse_drains_all_pages_before_done_and_chat_survives_long_trace(monkeypatch, tmp_path):
    from research_app import api

    local = Store(tmp_path)
    monkeypatch.setattr(api, "store", local)
    rid = local.create("long research")
    for i in range(505):
        local.emit(rid, "tool", str(i))
    local.emit(rid, "chat", "question", role="user", text="比较回撤")
    local.update(rid, status="complete")
    client = TestClient(api.app)
    result = client.get(f"/api/runs/{rid}/events").text
    assert result.count("id: ") == 506
    assert result.count("event: done") == 1
    last = local.events(rid, after=505)[0]["seq"]
    resumed = client.get(f"/api/runs/{rid}/events", headers={"Last-Event-ID": str(last)}).text
    assert "id: " not in resumed
    assert client.get(f"/api/runs/{rid}/chat").json() == [{"role": "user", "text": "比较回撤"}]


def test_chat_passes_prior_turns_and_keeps_researches_isolated(monkeypatch, tmp_path, bundle):
    from research_app import agents, api

    local = Store(tmp_path)
    monkeypatch.setattr(api, "store", local)
    monkeypatch.setattr(api, "config", Settings(tmp_path, "model", "token", "https://example.com"))
    rid = local.create("chat")
    path = local.run_dir(rid) / "bundle.json"
    local.save_json(path, bundle.model_dump(mode="json"))
    local.update(rid, status="complete", bundle_path=str(path))
    foreign = local.create("other research")
    local.emit(foreign, "chat", "unrelated", role="user", text="不应该出现在上下文")
    contexts = []

    class Runtime:
        def __init__(self, *args):
            pass

        async def explain(self, saved, question, selected, history=None):
            contexts.append(history)
            return "依据当前回撤指标回答"

    monkeypatch.setattr(agents, "AgentRuntime", Runtime)
    client = TestClient(api.app)
    assert client.post(f"/api/runs/{rid}/chat", json={"message": "解释最大回撤"}).status_code == 200
    assert client.post(f"/api/runs/{rid}/chat", json={"message": "它在哪段区间？"}).status_code == 200
    assert contexts == [
        [],
        [
            {"role": "user", "text": "解释最大回撤"},
            {"role": "assistant", "text": "依据当前回撤指标回答"},
        ],
    ]
